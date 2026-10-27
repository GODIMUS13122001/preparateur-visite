"""Écritures dans le CRM Excel, avec sauvegarde horodatée et contrôle après écriture.

  python scripts/ecrire_crm.py --brief-prepare R001
      -> Rendez-vous : « Brief préparé » = Oui
  python scripts/ecrire_crm.py --compte-rendu R001 --fichier sorties/comptes-rendus/R001_cr.json
      -> Rendez-vous : « Compte rendu de visite » rempli, Statut = Réalisé
      -> Activités : nouvelle ligne (ID suivant, formules recopiées, tableau tbl_Activites agrandi)
      -> Transactions : « Prochaine étape » (et la phase, si elle a été validée)
     À lancer UNIQUEMENT après la validation explicite du commercial.
Contraintes : openpyxl sans data_only, écriture uniquement dans les colonnes de saisie,
formules, tableaux et listes de validation conservés, fichier relu après écriture.
"""
from __future__ import annotations

import argparse
import json
import shutil
import sys
from datetime import datetime

import openpyxl

from crm_lib import DOSSIER_DATA, chemin_crm


def empreinte(wb) -> dict:
    """Ce qui ne doit pas changer lors d'une écriture : formules, tableaux, validations."""
    res = {}
    for ws in wb.worksheets:
        formules = sum(1 for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
        tables = sorted(f"{t.name}:{t.ref}" for t in ws.tables.values())
        dv = sorted(f"{d.sqref}|{d.formula1}" for d in ws.data_validations.dataValidation)
        res[ws.title] = {"formules": formules, "tables": tables, "validations": dv}
    res["noms_definis"] = sorted(wb.defined_names.keys())
    return res


def colonne(ws, entete: str) -> int:
    for cell in ws[1]:
        if cell.value and str(cell.value).strip() == entete:
            return cell.column
    raise SystemExit(f"Colonne « {entete} » introuvable dans l'onglet {ws.title}")


def ligne_rdv(ws, rid: str) -> int:
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, 1).value).strip().upper() == rid.upper():
            return r
    raise SystemExit(f"Rendez-vous {rid} introuvable")


def sauvegarder(chemin) -> str:
    dossier = DOSSIER_DATA / "sauvegardes"
    dossier.mkdir(parents=True, exist_ok=True)
    cible = dossier / f"{chemin.stem}_{datetime.now():%Y%m%d-%H%M%S}{chemin.suffix}"
    shutil.copy2(chemin, cible)
    return cible.name


def brief_prepare(rid: str) -> dict:
    chemin = chemin_crm()
    wb = openpyxl.load_workbook(chemin)
    ws = wb["Rendez-vous"]
    col, row = colonne(ws, "Brief préparé"), ligne_rdv(ws, rid)
    avant = ws.cell(row, col).value
    if avant == "Oui":
        return {"statut": "deja_fait", "rdv": rid, "message": "La case « Brief préparé » était déjà à Oui : aucune écriture."}
    reference = empreinte(wb)
    sauvegarde = sauvegarder(chemin)
    ws.cell(row, col).value = "Oui"
    wb.calculation.fullCalcOnLoad = True  # openpyxl efface les valeurs en cache : Excel recalcule à l'ouverture
    wb.save(chemin)

    # Relecture et contrôles
    relu = openpyxl.load_workbook(chemin)
    valeur = relu["Rendez-vous"].cell(row, col).value
    apres = empreinte(relu)
    problemes = [o for o in reference if reference[o] != apres.get(o)]
    if valeur != "Oui" or problemes:
        shutil.copy2(DOSSIER_DATA / "sauvegardes" / sauvegarde, chemin)
        return {"statut": "erreur", "rdv": rid, "message": "Contrôle après écriture en échec : fichier restauré.",
                "valeur_relue": valeur, "elements_modifies": problemes}
    return {"statut": "ok", "rdv": rid, "cellule": f"Rendez-vous!{ws.cell(row, col).coordinate}",
            "avant": avant, "apres": valeur, "sauvegarde": sauvegarde,
            "controle": "formules, tableaux et listes de validation identiques"}


RESULTATS = ("Positif", "Neutre", "À relancer", "Négatif")


def texte_brut(md: str) -> str:
    """Le compte rendu tel qu'on le lit dans une cellule Excel : sans la mise en forme markdown."""
    import re
    lignes = []
    for l in md.splitlines():
        l = re.sub(r"\*\*([^*]+)\*\*", r"\1", l).rstrip()
        if l.startswith("#"):
            l = l.lstrip("#").strip().upper() if l.startswith("## ") else l.lstrip("#").strip()
        elif l.startswith(("- ", "* ")):
            l = "• " + l[2:]
        lignes.append(l)
    return "\n".join(lignes).strip()


def compte_rendu(rid: str, fichier: str) -> dict:
    from copy import copy
    from datetime import date as _date, timedelta
    from openpyxl.formula.translate import Translator

    cr = json.loads(open(fichier, encoding="utf-8").read())
    maj = cr.get("maj_crm") or {}
    chemin = chemin_crm()
    wb = openpyxl.load_workbook(chemin)
    reference = empreinte(wb)

    rv = wb["Rendez-vous"]
    row = ligne_rdv(rv, rid)
    c_cr, c_statut = colonne(rv, "Compte rendu de visite"), colonne(rv, "Statut")
    if rv.cell(row, c_statut).value == "Réalisé" and rv.cell(row, c_cr).value:
        return {"statut": "deja_fait", "rdv": rid, "message": "Ce compte rendu est déjà enregistré dans le CRM : aucune écriture."}
    val = lambda entete: rv.cell(row, colonne(rv, entete)).value
    texte = texte_brut(cr["texte"])

    # --- Activités : nouvelle ligne, en recopiant formules et mise en forme de la ligne précédente
    act = wb["Activités"]
    derniere = max(r for r in range(2, act.max_row + 1) if act.cell(r, 1).value)
    nouvelle = derniere + 1
    numeros = [int(str(act.cell(r, 1).value)[1:]) for r in range(2, nouvelle) if str(act.cell(r, 1).value or "").startswith("A")]
    aid = f"A{max(numeros) + 1:04d}"
    debut, fin = val("Heure de début"), val("Heure de fin")
    duree = None
    if debut and fin:
        duree = int((datetime.combine(_date.today(), fin) - datetime.combine(_date.today(), debut)).total_seconds() // 60)
    resultat = maj.get("resultat") if maj.get("resultat") in RESULTATS else None
    echeance = None
    if maj.get("echeance"):
        try:
            echeance = datetime.fromisoformat(str(maj["echeance"])[:10])
        except ValueError:
            echeance = None
    saisie = {
        "ID activité": aid, "Date": val("Date"), "Type d'activité": "Réunion visio" if val("Type de rendez-vous") == "Visio" else "Visite terrain",
        "ID entreprise": val("ID entreprise"), "ID contact": val("ID contact"), "ID commercial": val("ID commercial"),
        "Objet": val("Objet du rendez-vous"), "Compte rendu / contenu": texte, "Résultat": resultat,
        "ID transaction associée": val("ID transaction associée"), "Prochaine action": maj.get("prochaine_action"),
        "Échéance prochaine action": echeance, "Durée (min)": duree, "Autres participants client": None,
    }
    entetes = {str(act.cell(1, c).value).strip(): c for c in range(1, act.max_column + 1) if act.cell(1, c).value}
    formules = 0
    for nom, c in entetes.items():
        source, cible = act.cell(derniere, c), act.cell(nouvelle, c)
        cible._style = copy(source._style)
        if isinstance(source.value, str) and source.value.startswith("="):
            cible.value = Translator(source.value, origin=source.coordinate).translate_formula(cible.coordinate)
            formules += 1
        elif nom in saisie:
            cible.value = saisie[nom]
    t = act.tables["tbl_Activites"]
    ancienne_ref = t.ref
    t.ref = f"{t.ref.split(':')[0]}:{openpyxl.utils.get_column_letter(act.max_column)}{nouvelle}"
    if t.autoFilter is not None:
        t.autoFilter.ref = t.ref

    # --- Rendez-vous : compte rendu et statut
    rv.cell(row, c_cr).value = texte
    rv.cell(row, c_statut).value = "Réalisé"

    # --- Transactions : prochaine étape (et phase si proposée puis validée)
    tid = val("ID transaction associée")
    maj_affaire = []
    if tid:
        tr = wb["Transactions"]
        tl = next((r for r in range(2, tr.max_row + 1) if tr.cell(r, 1).value == tid), None)
        if tl:
            if maj.get("prochaine_etape_affaire"):
                tr.cell(tl, colonne(tr, "Prochaine étape")).value = maj["prochaine_etape_affaire"]
                maj_affaire.append("prochaine étape")
            phases = [wb["Paramètres"].cell(r, 1).value for r in range(7, 14)]
            if maj.get("phase") in phases and maj["phase"] != tr.cell(tl, colonne(tr, "Phase de la transaction")).value:
                tr.cell(tl, colonne(tr, "Phase de la transaction")).value = maj["phase"]
                maj_affaire.append("phase")

    sauvegarde = sauvegarder(chemin)
    wb.calculation.fullCalcOnLoad = True
    wb.save(chemin)

    # --- Relecture et contrôles : seuls changements attendus = la ligne ajoutée et le tableau agrandi
    relu = openpyxl.load_workbook(chemin)
    apres = empreinte(relu)
    attendu = json.loads(json.dumps(reference))
    attendu["Activités"]["formules"] += formules
    attendu["Activités"]["tables"] = sorted(x.replace(f"tbl_Activites:{ancienne_ref}", f"tbl_Activites:{t.ref}") for x in attendu["Activités"]["tables"])
    problemes = [o for o in attendu if attendu[o] != apres.get(o)]
    ok = (relu["Rendez-vous"].cell(row, c_statut).value == "Réalisé" and relu["Activités"].cell(nouvelle, 1).value == aid)
    if problemes or not ok:
        shutil.copy2(DOSSIER_DATA / "sauvegardes" / sauvegarde, chemin)
        return {"statut": "erreur", "rdv": rid, "message": "Contrôle après écriture en échec : fichier restauré.", "elements_modifies": problemes}
    return {"statut": "ok", "rdv": rid, "activite": aid, "ligne": nouvelle, "formules_recopiees": formules,
            "tableau": f"tbl_Activites {ancienne_ref} -> {t.ref}", "affaire": maj_affaire, "sauvegarde": sauvegarde,
            "controle": "formules, tableaux et listes de validation conformes"}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--brief-prepare", metavar="RDV")
    g.add_argument("--compte-rendu", metavar="RDV")
    ap.add_argument("--fichier", help="compte rendu validé (JSON), avec --compte-rendu")
    args = ap.parse_args()
    if args.compte_rendu:
        if not args.fichier:
            raise SystemExit("--fichier est obligatoire avec --compte-rendu")
        res = compte_rendu(args.compte_rendu.upper(), args.fichier)
    else:
        res = brief_prepare(args.brief_prepare)
    print(json.dumps(res, ensure_ascii=False, indent=2))
    return 0 if res["statut"] in ("ok", "deja_fait") else 1


if __name__ == "__main__":
    sys.exit(main())
