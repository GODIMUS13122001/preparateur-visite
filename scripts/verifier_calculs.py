"""Compare les indicateurs recalculés par crm_lib aux valeurs en cache d'Excel.

Ne fonctionne que sur un fichier jamais réenregistré par openpyxl (cache présent).
Usage : python scripts/verifier_calculs.py [--fichier data/xxx.xlsx]
"""
import argparse
import sys

import openpyxl

from crm_lib import CRM, en_date


def egal(a, b):
    if isinstance(a, (int, float)) or isinstance(b, (int, float)):
        return abs((a or 0) - (b or 0)) < 0.01
    return (en_date(a) if a else None) == (en_date(b) if b else None) or a == b


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--fichier")
    args = ap.parse_args()
    crm = CRM(args.fichier)
    cache = openpyxl.load_workbook(crm.chemin, data_only=True)
    if cache["Entreprises"]["AA2"].value is None:
        print("Pas de valeurs en cache : fichier déjà réenregistré par openpyxl, vérification impossible.")
        return 2
    ecarts, controles = [], 0

    def check(lieu, calc, attendu):
        nonlocal controles
        controles += 1
        if not egal(calc, attendu):
            ecarts.append(f"{lieu}: recalculé={calc!r} Excel={attendu!r}")

    ws = cache["Commandes"]
    for i, c in enumerate(crm.lignes("Commandes"), start=2):
        check(f"Commandes!K{i}", crm.montant_commande(c), ws[f"K{i}"].value)
        r = ws[f"N{i}"].value
        check(f"Commandes!N{i}", crm.retard_commande(c), None if r == "" else r)
    ws = cache["Entreprises"]
    for i, e in enumerate(crm.lignes("Entreprises"), start=2):
        ind = crm.indicateurs_entreprise(e["ID entreprise"])
        for col, cle in [("AA", "ca_12_mois"), ("AB", "transactions_ouvertes"), ("AC", "pipeline_ouvert"),
                         ("AD", "tickets_ouverts"), ("AE", "date_derniere_activite"), ("AF", "prochain_rendez_vous")]:
            v = ws[f"{col}{i}"].value
            check(f"Entreprises!{col}{i}", ind[cle], None if v == "" else v)
    ws = cache["Contacts"]
    for i, c in enumerate(crm.lignes("Contacts"), start=2):
        v = ws[f"W{i}"].value
        check(f"Contacts!W{i}", crm.derniere_activite(cid=c["ID contact"]), None if v == "" else v)
        check(f"Contacts!X{i}", crm.interactions_12_mois(c["ID contact"]), ws[f"X{i}"].value)
    ws = cache["Transactions"]
    for i, t in enumerate(crm.lignes("Transactions"), start=2):
        info = crm.phase_info(t)
        check(f"Transactions!K{i}", info["probabilite"], ws[f"K{i}"].value)
        check(f"Transactions!N{i}", info["statut"], ws[f"N{i}"].value)
    for onglet, col_id, col_nom, fn in [("Activités", "D", "E", crm.nom_entreprise), ("Activités", "F", "G", crm.nom_contact),
                                         ("Rendez-vous", "H", "I", crm.nom_contact), ("Transactions", "C", "D", crm.nom_entreprise)]:
        ws = cache[onglet]
        for i in range(2, ws.max_row + 1):
            if ws[f"A{i}"].value:
                v = ws[f"{col_nom}{i}"].value
                check(f"{onglet}!{col_nom}{i}", fn(ws[f"{col_id}{i}"].value), None if v == "" else v)

    print(f"{controles} contrôles, {len(ecarts)} écart(s)")
    for e in ecarts[:30]:
        print("  ", e)
    return 1 if ecarts else 0


if __name__ == "__main__":
    sys.exit(main())
