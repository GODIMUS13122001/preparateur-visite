"""Lecture du CRM Excel Orvane, sans dépendre des colonnes calculées par formule.

Les valeurs en cache des formules disparaissent dès qu'openpyxl réenregistre le
fichier. Ce module ne lit donc que les colonnes de saisie (en-têtes bleu foncé)
et recalcule lui-même jointures et indicateurs à partir des identifiants.
La date de référence est Paramètres!B3 (date de l'export), jamais la date système.
"""
from __future__ import annotations

import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl

RACINE = Path(__import__("os").environ.get("PREPARATEUR_ESPACE") or Path(__file__).resolve().parent.parent)  # dossier des données
DOSSIER_DATA = RACINE / "data"
FICHIER_ACTIF = DOSSIER_DATA / ".crm_actif"
FICHIER_DEFAUT = "Export_CRM_Orvane_demo_20261009.xlsx"

GRIS_FORMULE = "FF5F6B7A"
ORANGE_AGENT = "FFC55A11"

ONGLETS = ["Commerciaux", "Entreprises", "Contacts", "Transactions", "Activités",
           "Rendez-vous", "Tickets", "Commandes", "Produits"]


def chemin_crm() -> Path:
    """Fichier CRM actif : celui choisi dans l'interface, sinon le fichier de démo."""
    if FICHIER_ACTIF.exists():
        nom = FICHIER_ACTIF.read_text(encoding="utf-8").strip()
        if nom and (DOSSIER_DATA / nom).exists():
            return DOSSIER_DATA / nom
    return DOSSIER_DATA / FICHIER_DEFAUT


def sans_accents(texte: str) -> str:
    texte = unicodedata.normalize("NFKD", texte or "")
    return "".join(c for c in texte if not unicodedata.combining(c)).lower().strip()


def slug(texte: str) -> str:
    s = sans_accents(texte)
    return "".join(c if c.isalnum() else "-" for c in s).strip("-").replace("--", "-")


def en_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def moins_mois(d: date, mois: int) -> date:
    """d moins N mois calendaires (jour borné à la fin du mois)."""
    a, m = d.year, d.month - mois
    while m <= 0:
        m += 12
        a -= 1
    for jour in (d.day, 30, 29, 28):
        try:
            return date(a, m, jour)
        except ValueError:
            continue
    raise ValueError(d)


@dataclass
class Onglet:
    nom: str
    entetes: list[str]
    saisie: list[str]          # colonnes de saisie (lues)
    calculees: list[str]       # colonnes calculées (ignorées à la lecture)
    agent: list[str]           # colonnes à remplir par l'agent
    lignes: list[dict] = field(default_factory=list)


def _lire_onglet(ws) -> Onglet:
    entetes, saisie, calculees, agent = [], [], [], []
    for cell in ws[1]:
        if cell.value is None:
            continue
        nom = str(cell.value).strip()
        entetes.append(nom)
        couleur = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type else None
        exemple = ws.cell(2, cell.column).value
        est_formule = isinstance(exemple, str) and exemple.startswith("=")
        if couleur == GRIS_FORMULE or est_formule:
            calculees.append(nom)
        else:
            saisie.append(nom)
            if couleur == ORANGE_AGENT:
                agent.append(nom)
    onglet = Onglet(ws.title, entetes, saisie, calculees, agent)
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        ligne = {}
        for nom, val in zip(entetes, row):
            if nom in saisie:
                ligne[nom] = val
        onglet.lignes.append(ligne)
    return onglet


class CRM:
    """Vue en lecture seule du CRM, avec indicateurs recalculés."""

    def __init__(self, chemin: Path | None = None):
        self.chemin = Path(chemin) if chemin else chemin_crm()
        wb = openpyxl.load_workbook(self.chemin)  # sans data_only : on ne lit jamais le cache
        self.onglets = {nom: _lire_onglet(wb[nom]) for nom in ONGLETS}
        p = wb["Paramètres"]
        self.date_ref: date = en_date(p["B3"].value)
        if self.date_ref is None:
            raise ValueError("Paramètres!B3 (date de l'export) est vide ou n'est pas une date")
        self.phases = {}
        for r in range(7, 30):
            phase = p.cell(r, 1).value
            if phase:
                self.phases[phase] = {"probabilite": p.cell(r, 2).value, "statut": p.cell(r, 3).value}
        self._index()

    # -- index et jointures ---------------------------------------------------
    def _index(self):
        self.commerciaux = {l["ID commercial"]: l for l in self.lignes("Commerciaux")}
        self.entreprises = {l["ID entreprise"]: l for l in self.lignes("Entreprises")}
        self.contacts = {l["ID contact"]: l for l in self.lignes("Contacts")}
        self.produits = {l["ID produit"]: l for l in self.lignes("Produits")}
        self.transactions = {l["ID transaction"]: l for l in self.lignes("Transactions")}
        self.rdv = {l["ID rendez-vous"]: l for l in self.lignes("Rendez-vous")}

    def lignes(self, onglet: str) -> list[dict]:
        return self.onglets[onglet].lignes

    def nom_commercial(self, uid):
        c = self.commerciaux.get(uid)
        return f"{c['Prénom']} {c['Nom']}" if c else None

    def nom_contact(self, cid):
        c = self.contacts.get(cid)
        return f"{c['Prénom']} {c['Nom']}" if c else None

    def nom_entreprise(self, eid):
        e = self.entreprises.get(eid)
        return e["Nom de l'entreprise"] if e else None

    def nom_produit(self, pid):
        p = self.produits.get(pid)
        return p["Nom du produit"] if p else None

    # -- indicateurs recalculés ----------------------------------------------
    def montant_commande(self, c) -> float | None:
        p = self.produits.get(c.get("ID produit"))
        if c.get("Quantité") in (None, "") or not p or p.get("Prix unitaire HT (€)") in (None, ""):
            return None
        return round(c["Quantité"] * p["Prix unitaire HT (€)"] * (1 - (c.get("Remise") or 0)), 2)

    @staticmethod
    def retard_commande(c) -> int | None:
        prevue, effective = en_date(c.get("Date de livraison prévue")), en_date(c.get("Date de livraison effective"))
        if effective is None or prevue is None:
            return None
        return max(0, (effective - prevue).days)

    def phase_info(self, t) -> dict:
        return self.phases.get(t.get("Phase de la transaction"), {"probabilite": None, "statut": None})

    def ca_12_mois(self, eid, decalage_jours: int = 0) -> float:
        fin = self.date_ref - timedelta(days=decalage_jours)
        debut = fin - timedelta(days=365)
        total = 0.0
        for c in self.lignes("Commandes"):
            d = en_date(c.get("Date de commande"))
            if c.get("ID entreprise") == eid and d and debut < d <= fin:
                total += self.montant_commande(c) or 0
        return round(total, 2)

    def derniere_activite(self, eid=None, cid=None) -> date | None:
        dates = [en_date(a["Date"]) for a in self.lignes("Activités")
                 if (eid is None or a.get("ID entreprise") == eid)
                 and (cid is None or a.get("ID contact") == cid)
                 and en_date(a["Date"]) and en_date(a["Date"]) <= self.date_ref]
        return max(dates) if dates else None

    def interactions_12_mois(self, cid) -> int:
        debut = self.date_ref - timedelta(days=365)
        return sum(1 for a in self.lignes("Activités")
                   if a.get("ID contact") == cid and en_date(a["Date"])
                   and debut < en_date(a["Date"]) <= self.date_ref)

    def indicateurs_entreprise(self, eid) -> dict:
        ouvertes = [t for t in self.lignes("Transactions")
                    if t.get("ID entreprise") == eid and self.phase_info(t)["statut"] == "Ouverte"]
        futurs = [en_date(r["Date"]) for r in self.lignes("Rendez-vous")
                  if r.get("ID entreprise") == eid and en_date(r["Date"]) and en_date(r["Date"]) >= self.date_ref]
        return {
            "ca_12_mois": self.ca_12_mois(eid),
            "transactions_ouvertes": len(ouvertes),
            "pipeline_ouvert": sum(t.get("Montant (€)") or 0 for t in ouvertes),
            "tickets_ouverts": sum(1 for k in self.lignes("Tickets")
                                   if k.get("ID entreprise") == eid and k.get("Statut") != "Fermé"),
            "date_derniere_activite": self.derniere_activite(eid=eid),
            "prochain_rendez_vous": min(futurs) if futurs else None,
        }

    def rdv_du_jour(self, jour: date, uid: str | None = None) -> list[dict]:
        res = [r for r in self.lignes("Rendez-vous") if en_date(r["Date"]) == jour
               and (uid is None or r.get("ID commercial") == uid)]
        return sorted(res, key=lambda r: str(r.get("Heure de début")))

    def prochain_rdv_entreprise(self, nom: str) -> dict | None:
        cible = sans_accents(nom)
        eids = [eid for eid, e in self.entreprises.items() if cible in sans_accents(e["Nom de l'entreprise"])]
        futurs = [r for r in self.lignes("Rendez-vous") if r.get("ID entreprise") in eids
                  and en_date(r["Date"]) and en_date(r["Date"]) >= self.date_ref
                  and r.get("Statut") not in ("Annulé", "Réalisé")]
        futurs.sort(key=lambda r: (en_date(r["Date"]), str(r.get("Heure de début"))))
        return futurs[0] if futurs else None
