"""Écriture dans le CRM de prospection (copie de sauvegarde horodatée avant chaque écriture).

Usage : .venv/bin/python scripts/ecrire_crm_prospection.py --brief-prepare RP003
Rendez-vous : « Brief préparé » = Oui. C'est la seule écriture de la préparation d'une fiche
(le commercial a demandé la fiche) ; les autres écritures viendront après validation d'un compte rendu.
"""
from __future__ import annotations

import argparse
import json
import shutil
from datetime import datetime

import openpyxl

from crm_lib import DOSSIER_DATA
from crm_prospection import CRMProspection, chemin_crm


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--brief-prepare", required=True, metavar="RDV")
    rid = ap.parse_args().brief_prepare.upper()
    chemin = chemin_crm()
    if chemin is None:
        raise SystemExit(json.dumps({"statut": "erreur", "message": "Aucun CRM de prospection déposé"}, ensure_ascii=False))
    wb = openpyxl.load_workbook(chemin)
    ws = wb["Rendez-vous"]
    entetes = [c.value for c in ws[1]]
    col_id, col = entetes.index("ID rendez-vous") + 1, entetes.index("Brief préparé") + 1
    ligne = next((r for r in range(2, ws.max_row + 1) if str(ws.cell(r, col_id).value).upper() == rid), None)
    if ligne is None:
        raise SystemExit(json.dumps({"statut": "introuvable", "message": f"{rid} absent du CRM de prospection"}, ensure_ascii=False))
    cellule = f"Rendez-vous!{openpyxl.utils.get_column_letter(col)}{ligne}"
    if ws.cell(ligne, col).value == "Oui":
        print(json.dumps({"statut": "deja_fait", "cellule": cellule}, ensure_ascii=False))
        return
    (DOSSIER_DATA / "sauvegardes").mkdir(exist_ok=True)
    sauvegarde = DOSSIER_DATA / "sauvegardes" / f"{chemin.stem}_{datetime.now():%Y%m%d-%H%M%S}.xlsx"
    shutil.copy2(chemin, sauvegarde)
    ws.cell(ligne, col).value = "Oui"
    wb.save(chemin)
    ok = CRMProspection(chemin).rdv[rid].get("Brief préparé") == "Oui"  # relecture
    print(json.dumps({"statut": "ok" if ok else "erreur", "cellule": cellule, "sauvegarde": sauvegarde.name}, ensure_ascii=False))


if __name__ == "__main__":
    main()
